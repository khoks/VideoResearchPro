"""PPTX rendering for decks (S-6.3.1), with python-pptx (MIT).

One editorial template drawn in code rather than shipped as a .potx: a
16:9 page in the brand's paper-and-ink palette (docs/branding.md), oxblood
accent rules, serif type. Five slide kinds: title, section, bullets, quote,
sources. Every content slide carries its sources twice: numbered in a footer
on the slide, and in full (title, channel, timestamp, link) in its notes.

Fonts: PowerPoint has no fallback list, so the deck names faces that exist
on most machines: Georgia for Latin text and Nirmala UI as the complex-script
face, which is what renders Devanagari quotes.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import datetime, timezone

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# docs/branding.md, light theme.
PAPER = RGBColor(0xFA, 0xF6, 0xF0)
SUNKEN = RGBColor(0xF0, 0xE9, 0xDD)
INK = RGBColor(0x1F, 0x1B, 0x16)
SECONDARY = RGBColor(0x5C, 0x54, 0x48)
MUTED = RGBColor(0x7D, 0x75, 0x6B)
OXBLOOD = RGBColor(0x7A, 0x2C, 0x1A)
GOLD = RGBColor(0xC7, 0x99, 0x45)

LATIN_FACE = "Georgia"
COMPLEX_FACE = "Nirmala UI"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.8)
SOURCES_PER_SLIDE = 8


@dataclass
class DeckSource:
    """One source document (a video), numbered by first use, with every
    moment the deck cites from it."""

    number: int
    title: str
    channel: str
    url: str  # the document without a timestamp, or its first cited moment
    moments: list[tuple[str, str]] = field(default_factory=list)  # (m:ss, link)

    def line(self) -> str:
        return f"{self.title} — {self.channel}" if self.channel else self.title


@dataclass
class DeckSlide:
    kind: str  # "section" | "bullets" | "quote"
    heading: str = ""
    bullets: list[str] = field(default_factory=list)
    quote: str = ""
    attribution: str = ""
    notes: str = ""
    sources: list[int] = field(default_factory=list)  # deck source numbers
    # source number -> the moments this slide cites from it, for the notes
    moments: dict[int, list[tuple[str, str]]] = field(default_factory=dict)


@dataclass
class Deck:
    title: str
    subtitle: str
    byline: str
    slides: list[DeckSlide]
    sources: list[DeckSource]
    disclosure: str = ""
    notes: list[str] = field(default_factory=list)  # processing notes for the last slide


def _font(run, size: float, color: RGBColor, *, bold: bool = False, italic: bool = False) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = LATIN_FACE
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", COMPLEX_FACE if tag == "a:cs" else LATIN_FACE)


def _background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PAPER


def _rule(slide, left, top, width, height=Inches(0.06), color=OXBLOOD) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False


def _text(slide, left, top, width, height, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Emu(0))
    return tf


def _paragraph(tf, text: str, size: float, color: RGBColor, *, first: bool = False, bold=False, italic=False,
               align=PP_ALIGN.LEFT, space_after: float = 0, url: str | None = None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    if space_after:
        p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    _font(run, size, color, bold=bold, italic=italic)
    if url:
        run.hyperlink.address = url
    return p


def _footer(slide, deck: Deck, page: int, sources: list[int] | None = None) -> None:
    tf = _text(slide, MARGIN, SLIDE_H - Inches(0.55), SLIDE_W - 2 * MARGIN, Inches(0.3), anchor=MSO_ANCHOR.MIDDLE)
    label = f"{deck.title} · {page}"
    if sources:
        label += "   ·   Sources " + ", ".join(str(n) for n in sources)
    _paragraph(tf, label, 10, MUTED, first=True)


def _notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def source_notes(deck: Deck, item: DeckSlide) -> str:
    """The slide's sources in full: each document once, then the exact
    moments this slide draws on, each with its link."""
    by_number = {s.number: s for s in deck.sources}
    lines = ["Sources:"]
    for n in item.sources:
        s = by_number.get(n)
        if not s:
            continue
        moments = item.moments.get(n) or []
        if moments:
            lines.append(f"[{n}] {s.line()}: " + "; ".join(f"{clock} {url}" if clock else url for clock, url in moments))
        else:
            lines.append(f"[{n}] {s.line()}: {s.url}")
    return "\n".join(lines)


def _source_entry(tf, s: DeckSource, first: bool) -> None:
    """``n.  Title — Channel   3:23 · 6:51``, each timestamp a link (the title
    links when there is no timestamp)."""
    p = _paragraph(tf, f"{s.number}.  ", 15, MUTED, first=first, space_after=6)
    run = p.add_run()
    run.text = s.line()
    _font(run, 15, INK)
    timed = [(clock, url) for clock, url in s.moments if clock]
    if not timed:
        run.hyperlink.address = s.url
        _font(run, 15, OXBLOOD)
        return
    for i, (clock, url) in enumerate(timed):
        sep = p.add_run()
        sep.text = "   " if i == 0 else " · "
        _font(sep, 13, MUTED)
        link = p.add_run()
        link.text = clock
        _font(link, 13, OXBLOOD)
        link.hyperlink.address = url


def render(deck: Deck) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]
    width = SLIDE_W - 2 * MARGIN

    # Title.
    slide = prs.slides.add_slide(blank)
    _background(slide)
    _rule(slide, MARGIN, Inches(2.2), Inches(1.6))
    tf = _text(slide, MARGIN, Inches(2.45), width, Inches(1.9), anchor=MSO_ANCHOR.TOP)
    _paragraph(tf, deck.title, 44, INK, first=True, bold=True, space_after=10)
    if deck.subtitle:
        _paragraph(tf, deck.subtitle, 22, SECONDARY, italic=True)
    tf = _text(slide, MARGIN, Inches(5.6), width, Inches(1.0))
    _paragraph(tf, deck.byline, 14, MUTED, first=True, space_after=4)
    if deck.disclosure:
        _paragraph(tf, deck.disclosure, 12, MUTED, italic=True)
    _notes(slide, "\n\n".join(x for x in (deck.disclosure, *deck.notes) if x))

    page = 1
    for item in deck.slides:
        page += 1
        slide = prs.slides.add_slide(blank)
        _background(slide)
        if item.kind == "section":
            _rule(slide, MARGIN, Inches(3.2), Inches(1.2))
            tf = _text(slide, MARGIN, Inches(3.45), width, Inches(1.4))
            _paragraph(tf, item.heading, 36, OXBLOOD, first=True, bold=True)
            _footer(slide, deck, page)
            _notes(slide, item.notes)
            continue
        if item.kind == "quote":
            _rule(slide, MARGIN, Inches(1.4), Inches(0.08), height=Inches(3.6), color=GOLD)
            tf = _text(slide, MARGIN + Inches(0.5), Inches(1.4), width - Inches(0.5), Inches(3.6), anchor=MSO_ANCHOR.MIDDLE)
            _paragraph(tf, f"“{item.quote}”", 30, INK, first=True, italic=True, space_after=18)
            if item.attribution:
                _paragraph(tf, f"— {item.attribution}", 16, SECONDARY)
        else:
            tf = _text(slide, MARGIN, Inches(0.7), width, Inches(1.1), anchor=MSO_ANCHOR.BOTTOM)
            _paragraph(tf, item.heading, 32, INK, first=True, bold=True)
            _rule(slide, MARGIN, Inches(1.95), Inches(0.9))
            tf = _text(slide, MARGIN, Inches(2.3), width, Inches(4.4))
            for i, bullet in enumerate(item.bullets):
                _paragraph(tf, f"•  {bullet}", 22, INK, first=(i == 0), space_after=12)
        _footer(slide, deck, page, item.sources)
        _notes(slide, "\n\n".join(x for x in (item.notes, source_notes(deck, item)) if x))

    # Sources, as many slides as it takes.
    chunks = [deck.sources[i : i + SOURCES_PER_SLIDE] for i in range(0, len(deck.sources), SOURCES_PER_SLIDE)] or [[]]
    for c, chunk in enumerate(chunks, start=1):
        page += 1
        slide = prs.slides.add_slide(blank)
        _background(slide)
        heading = "Sources" if len(chunks) == 1 else f"Sources ({c} of {len(chunks)})"
        tf = _text(slide, MARGIN, Inches(0.7), width, Inches(0.9), anchor=MSO_ANCHOR.BOTTOM)
        _paragraph(tf, heading, 32, INK, first=True, bold=True)
        _rule(slide, MARGIN, Inches(1.75), Inches(0.9))
        tf = _text(slide, MARGIN, Inches(2.0), width, Inches(4.8))
        for i, s in enumerate(chunk):
            _source_entry(tf, s, first=(i == 0))
        if c == len(chunks) and deck.notes:
            for note in deck.notes:
                _paragraph(tf, note, 12, MUTED, italic=True, space_after=4)
        _footer(slide, deck, page)
        _notes(slide, "\n".join(f"[{s.number}] {s.line()}: " + "; ".join(f"{c} {u}".strip() for c, u in s.moments or [("", s.url)]) for s in chunk))

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def byline(runs: int, model: str, when: datetime | None = None) -> tuple[str, str]:
    when = when or datetime.now(timezone.utc)
    return (
        f"Compiled by Pratidhvani Author Studio from {runs} research run{'' if runs == 1 else 's'} · {when:%Y-%m-%d}",
        f"Slides planned by a model ({model}) from the reports; every point's sources are in the slide notes.",
    )
