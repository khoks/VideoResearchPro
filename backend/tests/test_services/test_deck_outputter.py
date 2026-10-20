"""S-6.3.1 — slide decks, with the model patched.

The acceptance: a user picks runs and gets a ``.pptx`` that opens, whose
every content slide has at least one citation in its notes, and whose sources
slide lists every cited source with a link. Underneath that: a slide citing
nothing (or only sources the deck cannot link) is dropped and disclosed,
sources are numbered by first use and recorded as ``output_citations``, and
the model only ever sees markers for resolvable sources.
"""
from __future__ import annotations

import io
import json
import re

import pytest
from langchain_core.messages import AIMessage
from pptx import Presentation

from app.agents.prompts import deck_prompts
from app.agents.prompts.shared import ENGLISH_OUTPUT_CONTRACT, VISUAL_ANNOTATION_CONTRACT
from app.models.output import Output
from app.models.output_citation import OutputCitation
from app.services import auth_service, output_artifacts, output_service
from app.services.outputters import book_cohesion
from tests.test_services.test_book_endnotes import _smoke_run

MARKER = re.compile(r"\[S(\d+)\]")


class FakeModel:
    def __init__(self, answer):
        self.answer = answer
        self.prompts: list[str] = []

    def __call__(self, use_case, system, user, max_tokens):
        assert use_case == "author_deck_outline"
        self.prompts.append(user)
        if isinstance(self.answer, Exception):
            raise self.answer
        answer = self.answer(user) if callable(self.answer) else self.answer
        return AIMessage(content=json.dumps(answer), usage_metadata={"input_tokens": 9_000, "output_tokens": 3_000, "total_tokens": 12_000})


def _plan_from(prompt: str, *, uncited: int = 1, foreign: bool = True) -> dict:
    """A plausible plan that cites the markers the prompt actually offered."""
    markers = sorted({int(n) for n in MARKER.findall(prompt)})
    assert len(markers) >= 10, "the fixture report offers enough resolvable sources"
    slides: list[dict] = [{"type": "section", "heading": "Why it matters"}]
    for i in range(9):
        slides.append({
            "type": "bullets", "heading": f"Point {i + 1} [S{markers[i]}]",
            "bullets": [f"A specific finding {i} [S{markers[i]}] with a [link](https://x.example)", "Another detail[^3]"],
            "speaker_notes": f"Say this about point {i + 1}.",
            "citations": [markers[i], f"S{markers[i + 1]}"] + ([99_999] if foreign else []),
        })
    slides.append({"type": "quote", "quote": "Smoking is the leading cause", "attribution": "A pulmonologist",
                   "speaker_notes": "Exact words matter here.", "citations": [markers[0]]})
    for _ in range(uncited):
        slides.append({"type": "bullets", "heading": "Unsourced", "bullets": ["A claim with no marker"], "speaker_notes": "x", "citations": []})
    slides.append({"type": "section", "heading": "Trailing section with nothing after it"})
    return {"title": "Smoke, pills and lungs", "subtitle": "What three videos say", "slides": slides}


@pytest.fixture
def user_and_run(db, tmp_path, monkeypatch):
    import app.services.outputters  # noqa: F401

    monkeypatch.setattr("app.config.settings.OUTPUTS_DIR", str(tmp_path / "outputs"))
    user = auth_service.create_user(db, email="deck@x.com", password="pw" * 6)
    return user, _smoke_run(db, user.id, tmp_path, topic="Smoke run")


def _generate(db, user, job_ids, model):
    row = output_service.create_output(db, user_id=user.id, kind="deck", title="My deck", source_ids=job_ids)
    out = output_service.run_generation(db, user, row)
    return out


def test_a_deck_opens_and_every_content_slide_cites_a_linked_source(db, user_and_run, monkeypatch):
    user, run = user_and_run
    model = FakeModel(lambda prompt: _plan_from(prompt))
    monkeypatch.setattr(book_cohesion, "_default_invoke", model)
    out = _generate(db, user, [run.id], model)
    assert out.status == "completed", out.error_message

    # The model saw the contracts and markers, never raw links or footnotes.
    assert "](http" not in model.prompts[0] and "[^" not in model.prompts[0]
    assert VISUAL_ANNOTATION_CONTRACT in deck_prompts.DECK_SYSTEM and ENGLISH_OUTPUT_CONTRACT in deck_prompts.DECK_SYSTEM

    assert output_artifacts.available_formats(out) == ["md", "pptx"]
    prs = Presentation(io.BytesIO(output_artifacts.artifact_file(out, "pptx").read_bytes()))
    slides = list(prs.slides)
    texts = ["\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in slides]
    notes = [s.notes_slide.notes_text_frame.text if s.has_notes_slide else "" for s in slides]

    assert "Smoke, pills and lungs" in texts[0] and "gpt-5.5" in texts[0], "title slide discloses the model"
    sources_at = [i for i, t in enumerate(texts) if t.startswith("Sources")]
    assert sources_at and sources_at[-1] == len(slides) - 1
    content = [i for i in range(1, sources_at[0]) if "Why it matters" not in texts[i]]
    assert 3 <= len(content) <= 15
    for i in content:
        assert "Sources:" in notes[i] and "https://www.youtube.com/watch?v=" in notes[i], f"slide {i + 1} notes carry its sources"
        assert "[S" not in texts[i] and "http" not in texts[i] and "[^" not in texts[i], f"slide {i + 1} text is clean"

    links = [r.hyperlink.address for i in sources_at for sh in slides[i].shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs for r in p.runs if r.hyperlink.address]
    rows = db.query(OutputCitation).filter(OutputCitation.output_id == out.id).order_by(OutputCitation.position).all()
    assert len(rows) >= 2 and all(r.url and r.url.startswith("https://") for r in rows)
    assert sorted(set(links)) == sorted({r.url for r in rows}), "every cited moment is linked from the sources slide"
    documents = {r.document_id for r in rows}
    entries = [p.text for i in sources_at for sh in slides[i].shapes if sh.has_text_frame
               for p in sh.text_frame.paragraphs if re.match(r"^\d+\.", p.text)]
    assert len(entries) == len(documents), "one entry per source video, however many moments it gives"
    for entry in entries:
        clocks = [int(m) * 60 + int(s) for m, s in re.findall(r"(\d+):(\d\d)", entry)]
        assert clocks == sorted(clocks), f"moments in time order: {entry}"

    # Disclosure of what was left out, on the deck and in its Markdown.
    assert "1 planned slide cited no source and was left out." in texts[-1]
    assert "Trailing section" not in "\n".join(texts)
    md = out.content_text
    assert md.startswith("# Smoke, pills and lungs") and "## Sources" in md
    assert len(re.findall(r"^\[\^\d+\]: ", md, flags=re.M)) == len(documents)
    usage = json.loads(out.usage_json)
    assert usage["calls"][0]["use_case"] == "author_deck_outline" and usage["cost_usd"] > 0


def test_sources_are_numbered_by_first_use(db, user_and_run, monkeypatch):
    user, run = user_and_run
    monkeypatch.setattr(book_cohesion, "_default_invoke", FakeModel(lambda p: _plan_from(p, uncited=0, foreign=False)))
    out = _generate(db, user, [run.id], None)
    prs = Presentation(io.BytesIO(output_artifacts.artifact_file(out, "pptx").read_bytes()))
    footers = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame and "Sources " in sh.text_frame.text and "·" in sh.text_frame.text]
    first = [int(n) for f in footers for n in re.findall(r"\d+", f.split("Sources ")[1])]
    seen: list[int] = []
    for n in first:
        if n not in seen:
            seen.append(n)
    assert seen == list(range(1, len(seen) + 1))


def test_a_plan_with_no_cited_slide_fails_the_output_and_says_why(db, user_and_run, monkeypatch):
    user, run = user_and_run
    plan = {"title": "T", "slides": [{"type": "bullets", "heading": "H", "bullets": ["b"], "citations": [12345]}]}
    monkeypatch.setattr(book_cohesion, "_default_invoke", FakeModel(plan))
    out = _generate(db, user, [run.id], None)
    assert out.status == "failed" and "no slide citing a source" in out.error_message
    assert output_artifacts.available_formats(out) == []


def test_a_failed_outline_call_fails_the_output(db, user_and_run, monkeypatch):
    user, run = user_and_run
    monkeypatch.setattr(book_cohesion, "_default_invoke", FakeModel(RuntimeError("provider down")))
    out = _generate(db, user, [run.id], None)
    assert out.status == "failed" and "provider down" in out.error_message


def test_another_users_run_is_not_material(db, user_and_run, tmp_path, monkeypatch):
    user, _run = user_and_run
    other = auth_service.create_user(db, email="deck-other@x.com", password="pw" * 6)
    theirs = _smoke_run(db, other.id, tmp_path, topic="Not yours")
    monkeypatch.setattr(book_cohesion, "_default_invoke", lambda *a: pytest.fail("no model call for nothing"))
    out = _generate(db, user, [theirs.id], None)
    assert out.status == "failed" and "None of the chosen runs" in out.error_message


def test_estimate_and_kinds(unauthenticated_client, db, user_and_run, monkeypatch):
    user, run = user_and_run
    user.tier = "pro"
    db.commit()
    token, _ = auth_service.create_access_token(user.id)
    headers = {"Authorization": f"Bearer {token}"}
    assert "deck" in unauthenticated_client.get("/api/v1/author/kinds", headers=headers).json()["available"]
    monkeypatch.setattr(book_cohesion, "_default_invoke", lambda *a: pytest.fail("estimate must not call a model"))
    est = unauthenticated_client.post("/api/v1/author/outputs/estimate", headers=headers,
                                      json={"kind": "deck", "title": "D", "source_ids": [run.id]}).json()
    assert [c["use_case"] for c in est["calls"]] == ["author_deck_outline"] and est["cost_usd"] > 0
    assert db.query(Output).filter(Output.kind == "deck").count() == 0
